"""Payroll export, import reconciliation, posting, and sign-off."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import uuid
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

from app.core.db import transaction
from app.core.errors import Conflict, Forbidden, NotFound
from app.core.time import payroll_period
from app.domains.audit.service import append_audit
from app.domains.auth.service import CurrentSession
from app.domains.repayments.service import post_repayment


@dataclass(frozen=True)
class ImportReport:
    file_sha256: str
    file_total_kobo: int
    rows: list[dict[str, Any]]
    errors: list[dict[str, Any]]


_CSV_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_csv_row(
    values: dict[str, Any],
    text_fields: set[str],
    max_lengths: dict[str, int] | None = None,
) -> dict[str, Any]:
    sanitized = dict(values)
    limits = max_lengths or {}
    for field in text_fields:
        value = sanitized.get(field)
        if not isinstance(value, str):
            continue
        if value.startswith(_CSV_FORMULA_PREFIXES):
            value = "'" + value
        if field in limits:
            value = value[: limits[field]]
        sanitized[field] = value
    return sanitized


def _role(current: CurrentSession, allowed: set[str]) -> str:
    role = next((item for item in current.roles if item in allowed), None)
    if role is None:
        raise Forbidden("This role cannot manage payroll reconciliation")
    return role


def period_for(moment: datetime) -> tuple[int, int]:
    return payroll_period(moment)


def _parse_money(value: object) -> int:
    text_value = str(value).strip().replace(",", "")
    if not text_value or text_value.startswith("-") or "." in text_value:
        raise ValueError("amount must be a non-negative whole-kobo integer")
    return int(text_value)


def _read_rows(
    content: bytes, filename: str, profile: dict[str, Any]
) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    try:
        if filename.lower().endswith(".xlsx"):
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            values = list(sheet.iter_rows(values_only=True))
            header = [str(value or "").strip() for value in values[int(profile["header_row"]) - 1]]
            records = [
                dict(zip(header, [str(value or "").strip() for value in row], strict=False))
                for row in values[int(profile["header_row"]) :]
            ]
        else:
            decoded = content.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(decoded)))
            header = rows[int(profile["header_row"]) - 1]
            records = [
                dict(zip(header, row, strict=False)) for row in rows[int(profile["header_row"]) :]
            ]
    except Exception as exc:
        return [], [{"row": 0, "error": f"Unreadable payroll file: {exc}"}]
    needed = {profile["staff_no_column"], profile["amount_column"]}
    missing = needed - set(header)
    if missing:
        return [], [
            {
                "row": int(profile["header_row"]),
                "error": f"Missing columns: {', '.join(sorted(missing))}",
            }
        ]
    return records, []


async def create_profile(current: CurrentSession, values: dict[str, Any]) -> uuid.UUID:
    _role(current, {"admin"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        version = await session.scalar(
            text(
                "SELECT coalesce(max(version),0)+1 FROM payroll_import_profiles WHERE cooperative_id=:cid"
            ),
            {"cid": cid},
        )
        profile_id = uuid.uuid4()
        await session.execute(
            text("""INSERT INTO payroll_import_profiles(cooperative_id,id,version,staff_no_column,amount_column,name_column,header_row,created_by)
          VALUES(:cid,:id,:version,:staff,:amount,:name,:header,:actor)"""),
            {
                "cid": cid,
                "id": profile_id,
                "version": version,
                "staff": values["staff_no_column"],
                "amount": values["amount_column"],
                "name": values.get("name_column"),
                "header": values.get("header_row", 1),
                "actor": current.user_id,
            },
        )
    return profile_id


async def create_run(
    current: CurrentSession, profile_id: uuid.UUID, year: int, month: int
) -> uuid.UUID:
    _role(current, {"officer", "treasurer"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run_id = uuid.uuid4()
        await session.execute(
            text("""INSERT INTO payroll_runs(cooperative_id,id,profile_id,period_year,period_month,created_by)
          VALUES(:cid,:id,:profile,:year,:month,:actor)"""),
            {
                "cid": cid,
                "id": run_id,
                "profile": profile_id,
                "year": year,
                "month": month,
                "actor": current.user_id,
            },
        )
    return run_id


async def list_runs(current: CurrentSession, page: int, page_size: int) -> dict[str, Any]:
    _role(current, {"officer", "treasurer", "admin"})
    cid = current.cooperative_id
    offset = (page - 1) * page_size
    async with transaction(cid) as session:
        total = await session.scalar(
            text("SELECT count(*) FROM payroll_runs WHERE cooperative_id=:cid"),
            {"cid": cid},
        )
        rows = (
            (
                await session.execute(
                    text(
                        """SELECT r.id,r.period_year,r.period_month,r.status,
                          r.expected_total_kobo,r.file_total_kobo,r.created_by,r.posted_by,
                          r.created_at,r.imported_at,r.posted_at,r.signed_off_at,
                          coalesce(c.matched,0)::int AS matched,
                          coalesce(c.short,0)::int AS short,
                          coalesce(c.over,0)::int AS over,
                          coalesce(c.missing,0)::int AS missing,
                          coalesce(c.unknown_member,0)::int AS unknown_member,
                          coalesce(c.duplicate,0)::int AS duplicate
                          FROM payroll_runs r
                          LEFT JOIN (
                            SELECT cooperative_id,run_id,
                              count(*) FILTER (WHERE classification='matched') AS matched,
                              count(*) FILTER (WHERE classification='short') AS short,
                              count(*) FILTER (WHERE classification='over') AS over,
                              count(*) FILTER (WHERE classification='missing') AS missing,
                              count(*) FILTER (WHERE classification='unknown_member') AS unknown_member,
                              count(*) FILTER (WHERE classification='duplicate') AS duplicate
                            FROM payroll_items
                            WHERE cooperative_id=:cid
                            GROUP BY cooperative_id,run_id
                          ) c ON c.cooperative_id=r.cooperative_id AND c.run_id=r.id
                          WHERE r.cooperative_id=:cid
                          ORDER BY r.period_year DESC,r.period_month DESC,r.created_at DESC
                          LIMIT :limit OFFSET :offset"""
                    ),
                    {"cid": cid, "limit": page_size, "offset": offset},
                )
            )
            .mappings()
            .all()
        )
    items = []
    for row in rows:
        items.append(
            {
                "run_id": str(row["id"]),
                "period_year": int(row["period_year"]),
                "period_month": int(row["period_month"]),
                "status": str(row["status"]),
                "expected_total_kobo": int(row["expected_total_kobo"]),
                "file_total_kobo": (
                    int(row["file_total_kobo"]) if row["file_total_kobo"] is not None else None
                ),
                "counts": {
                    key: int(row[key])
                    for key in (
                        "matched",
                        "short",
                        "over",
                        "missing",
                        "unknown_member",
                        "duplicate",
                    )
                },
                "created_by": str(row["created_by"]),
                "posted_by": str(row["posted_by"]) if row["posted_by"] is not None else None,
                "created_at": row["created_at"],
                "imported_at": row["imported_at"],
                "posted_at": row["posted_at"],
                "signed_off_at": row["signed_off_at"],
            }
        )
    return {"items": items, "page": page, "page_size": page_size, "total": int(total or 0)}


async def get_run(current: CurrentSession, run_id: uuid.UUID) -> dict[str, Any]:
    _role(current, {"officer", "treasurer", "admin"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run = (
            (
                await session.execute(
                    text(
                        """SELECT id,period_year,period_month,status,expected_total_kobo,
                          file_total_kobo,bank_receipt_total_kobo,created_by,posted_by,
                          created_at,imported_at,posted_at,signed_off_at
                          FROM payroll_runs WHERE cooperative_id=:cid AND id=:id"""
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        if run is None:
            raise NotFound("Payroll run not found")
        totals = (
            (
                await session.execute(
                    text(
                        """SELECT count(*) FILTER (WHERE classification='matched') AS matched,
                          count(*) FILTER (WHERE classification='short') AS short,
                          count(*) FILTER (WHERE classification='over') AS over,
                          count(*) FILTER (WHERE classification='missing') AS missing,
                          count(*) FILTER (WHERE classification='unknown_member') AS unknown_member,
                          count(*) FILTER (WHERE classification='duplicate') AS duplicate,
                          coalesce(sum(expected_loan_kobo),0)::bigint AS loan_kobo,
                          coalesce(sum(expected_savings_kobo),0)::bigint AS savings_kobo,
                          coalesce(sum(expected_dues_kobo),0)::bigint AS dues_kobo
                          FROM payroll_items WHERE cooperative_id=:cid AND run_id=:id"""
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one()
        )
        exception_rows = (
            (
                await session.execute(
                    text(
                        """SELECT staff_no_raw,classification,actual_kobo,resolution_note
                          FROM payroll_items WHERE cooperative_id=:cid AND run_id=:id
                            AND classification<>'matched'
                          ORDER BY created_at,id"""
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .all()
        )

    bank_total = run["bank_receipt_total_kobo"]
    file_total = run["file_total_kobo"]
    exceptions = []
    for exception in exception_rows:
        safe = _sanitize_csv_row(
            dict(exception),
            {"staff_no_raw", "classification", "resolution_note"},
            {"resolution_note": 500},
        )
        exceptions.append(
            {
                "staff_no_raw": safe["staff_no_raw"],
                "classification": safe["classification"],
                "actual_kobo": (
                    int(exception["actual_kobo"]) if exception["actual_kobo"] is not None else None
                ),
                "resolution_note": safe["resolution_note"],
            }
        )
    return {
        "run_id": str(run["id"]),
        "period_year": int(run["period_year"]),
        "period_month": int(run["period_month"]),
        "status": str(run["status"]),
        "expected_total_kobo": int(run["expected_total_kobo"]),
        "file_total_kobo": int(file_total) if file_total is not None else None,
        "bank_receipt_total_kobo": int(bank_total) if bank_total is not None else None,
        "variance_kobo": (
            int(file_total) - int(bank_total)
            if file_total is not None and bank_total is not None
            else None
        ),
        "counts": {
            key: int(totals[key])
            for key in ("matched", "short", "over", "missing", "unknown_member", "duplicate")
        },
        "component_totals_kobo": {
            "loan": int(totals["loan_kobo"]),
            "savings": int(totals["savings_kobo"]),
            "dues": int(totals["dues_kobo"]),
        },
        "created_by": str(run["created_by"]),
        "posted_by": str(run["posted_by"]) if run["posted_by"] is not None else None,
        "created_at": run["created_at"],
        "imported_at": run["imported_at"],
        "posted_at": run["posted_at"],
        "signed_off_at": run["signed_off_at"],
        "exceptions": exceptions,
    }


async def export_expected(current: CurrentSession, run_id: uuid.UUID) -> bytes:
    role = _role(current, {"officer", "treasurer"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run = (
            (
                await session.execute(
                    text(
                        """SELECT r.*,p.staff_no_column,p.amount_column,p.name_column FROM payroll_runs r JOIN payroll_import_profiles p ON p.cooperative_id=r.cooperative_id AND p.id=r.profile_id WHERE r.cooperative_id=:cid AND r.id=:id FOR UPDATE"""
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        if run is None:
            raise NotFound("Payroll run not found")
        if run["status"] not in {"draft", "exported"}:
            raise Conflict("Payroll run cannot be exported")
        rows = (
            (
                await session.execute(
                    text("""SELECT m.id,m.staff_no,m.full_name,coalesce(sum(s.principal_kobo+s.interest_kobo),0)::bigint loan_kobo,
          coalesce(sum(s.savings_kobo),0)::bigint savings_kobo,coalesce(sum(s.dues_kobo),0)::bigint dues_kobo
          FROM loan_schedule s JOIN loans l ON l.cooperative_id=s.cooperative_id AND l.id=s.loan_id JOIN members m ON m.cooperative_id=l.cooperative_id AND m.id=l.member_id
          WHERE s.cooperative_id=:cid AND l.status='repaying' AND NOT l.deductions_paused
            AND extract(year FROM s.due_at AT TIME ZONE 'Africa/Lagos')=:year
            AND extract(month FROM s.due_at AT TIME ZONE 'Africa/Lagos')=:month
          GROUP BY m.id,m.staff_no,m.full_name ORDER BY m.staff_no"""),
                    {"cid": cid, "year": run["period_year"], "month": run["period_month"]},
                )
            )
            .mappings()
            .all()
        )
        output = io.StringIO(newline="")
        headers = [run["staff_no_column"], run["amount_column"]] + (
            [run["name_column"]] if run["name_column"] else []
        )
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        total = 0
        for row in rows:
            amount = int(row["loan_kobo"]) + int(row["savings_kobo"]) + int(row["dues_kobo"])
            total += amount
            value = {run["staff_no_column"]: row["staff_no"], run["amount_column"]: amount}
            if run["name_column"]:
                value[run["name_column"]] = row["full_name"]
            text_fields = {run["staff_no_column"]}
            if run["name_column"]:
                text_fields.add(run["name_column"])
            writer.writerow(_sanitize_csv_row(value, text_fields))
        await session.execute(
            text(
                "UPDATE payroll_runs SET expected_total_kobo=:total,status='exported' WHERE cooperative_id=:cid AND id=:id"
            ),
            {"cid": cid, "id": run_id, "total": total},
        )
        await append_audit(
            session,
            cooperative_id=cid,
            actor_user_id=current.user_id,
            role=role,
            action="payroll_expected_exported",
            entity_type="payroll_run",
            entity_id=str(run_id),
            summary={
                "period_year": int(run["period_year"]),
                "period_month": int(run["period_month"]),
            },
            channel="api",
            ip=None,
        )
    return output.getvalue().encode()


async def dry_run(
    current: CurrentSession, run_id: uuid.UUID, content: bytes, filename: str
) -> ImportReport:
    _role(current, {"officer", "treasurer"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run = (
            (
                await session.execute(
                    text(
                        """SELECT r.*,p.staff_no_column,p.amount_column,p.name_column,p.header_row FROM payroll_runs r JOIN payroll_import_profiles p ON p.cooperative_id=r.cooperative_id AND p.id=r.profile_id WHERE r.cooperative_id=:cid AND r.id=:id"""
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        if run is None:
            raise NotFound("Payroll run not found")
        records, errors = _read_rows(content, filename, dict(run))
        staff = [record.get(run["staff_no_column"], "").strip() for record in records]
        members = (
            (
                await session.execute(
                    text(
                        "SELECT id,staff_no FROM members WHERE cooperative_id=:cid AND staff_no=ANY(:staff)"
                    ),
                    {"cid": cid, "staff": staff or [""]},
                )
            )
            .mappings()
            .all()
        )
        member_by_staff = {row["staff_no"]: row["id"] for row in members}
        seen: set[str] = set()
        output: list[dict[str, Any]] = []
        total = 0
        for index, record in enumerate(records, start=int(run["header_row"]) + 1):
            staff_no = record.get(run["staff_no_column"], "").strip()
            try:
                amount = _parse_money(record.get(run["amount_column"], ""))
            except ValueError as exc:
                errors.append({"row": index, "error": str(exc)})
                continue
            classification = (
                "duplicate"
                if staff_no in seen
                else ("unknown_member" if staff_no not in member_by_staff else "pending")
            )
            seen.add(staff_no)
            total += amount
            output.append(
                {
                    "row": index,
                    "staff_no_raw": staff_no,
                    "member_id": str(member_by_staff[staff_no])
                    if staff_no in member_by_staff
                    else None,
                    "actual_kobo": amount,
                    "classification": classification,
                }
            )
        return ImportReport(hashlib.sha256(content).hexdigest(), total, output, errors)


async def import_file(
    current: CurrentSession, run_id: uuid.UUID, content: bytes, filename: str
) -> ImportReport:
    role = _role(current, {"officer", "treasurer"})
    report = await dry_run(current, run_id, content, filename)
    if report.errors:
        return report
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run = (
            (
                await session.execute(
                    text(
                        "SELECT * FROM payroll_runs WHERE cooperative_id=:cid AND id=:id FOR UPDATE"
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        if run is None or run["status"] not in {"draft", "exported"}:
            raise Conflict("Payroll run cannot be imported")
        digest = bytes.fromhex(report.file_sha256)
        duplicate = await session.scalar(
            text(
                "SELECT EXISTS(SELECT 1 FROM payroll_runs WHERE cooperative_id=:cid AND period_year=:year AND period_month=:month AND file_sha256=:hash)"
            ),
            {"cid": cid, "year": run["period_year"], "month": run["period_month"], "hash": digest},
        )
        if duplicate:
            raise Conflict("This file was already imported for the payroll period")
        expected = (
            (
                await session.execute(
                    text(
                        """SELECT m.id,m.staff_no,coalesce(sum(s.principal_kobo+s.interest_kobo),0)::bigint loan_kobo,coalesce(sum(s.savings_kobo),0)::bigint savings_kobo,coalesce(sum(s.dues_kobo),0)::bigint dues_kobo FROM loan_schedule s JOIN loans l ON l.cooperative_id=s.cooperative_id AND l.id=s.loan_id JOIN members m ON m.cooperative_id=l.cooperative_id AND m.id=l.member_id WHERE s.cooperative_id=:cid AND l.status='repaying' AND NOT l.deductions_paused AND extract(year FROM s.due_at AT TIME ZONE 'Africa/Lagos')=:year AND extract(month FROM s.due_at AT TIME ZONE 'Africa/Lagos')=:month GROUP BY m.id,m.staff_no"""
                    ),
                    {"cid": cid, "year": run["period_year"], "month": run["period_month"]},
                )
            )
            .mappings()
            .all()
        )
        by_staff = {row["staff_no"]: row for row in expected}
        imported_staff: set[str] = set()
        for item in report.rows:
            staff = item["staff_no_raw"]
            expected_row = by_staff.get(staff)
            actual = int(item["actual_kobo"])
            imported_staff.add(staff)
            if item["classification"] != "pending":
                classification = item["classification"]
            elif expected_row is None:
                classification = "unknown_member"
            else:
                expected_total = (
                    int(expected_row["loan_kobo"])
                    + int(expected_row["savings_kobo"])
                    + int(expected_row["dues_kobo"])
                )
                classification = (
                    "matched"
                    if actual == expected_total
                    else ("short" if actual < expected_total else "over")
                )
            await session.execute(
                text(
                    """INSERT INTO payroll_items(cooperative_id,run_id,member_id,staff_no_raw,expected_loan_kobo,expected_savings_kobo,expected_dues_kobo,actual_kobo,classification) VALUES(:cid,:run,:member,:staff,:loan,:savings,:dues,:actual,:classification)"""
                ),
                {
                    "cid": cid,
                    "run": run_id,
                    "member": expected_row["id"] if expected_row else None,
                    "staff": staff,
                    "loan": expected_row["loan_kobo"] if expected_row else 0,
                    "savings": expected_row["savings_kobo"] if expected_row else 0,
                    "dues": expected_row["dues_kobo"] if expected_row else 0,
                    "actual": actual,
                    "classification": classification,
                },
            )
        for staff, row in by_staff.items():
            if staff not in imported_staff:
                await session.execute(
                    text(
                        """INSERT INTO payroll_items(cooperative_id,run_id,member_id,staff_no_raw,expected_loan_kobo,expected_savings_kobo,expected_dues_kobo,classification) VALUES(:cid,:run,:member,:staff,:loan,:savings,:dues,'missing')"""
                    ),
                    {
                        "cid": cid,
                        "run": run_id,
                        "member": row["id"],
                        "staff": staff,
                        "loan": row["loan_kobo"],
                        "savings": row["savings_kobo"],
                        "dues": row["dues_kobo"],
                    },
                )
        expected_total_kobo = sum(
            int(row["loan_kobo"]) + int(row["savings_kobo"]) + int(row["dues_kobo"])
            for row in by_staff.values()
        )
        await session.execute(
            text(
                "UPDATE payroll_runs SET status='imported',expected_total_kobo=:expected,file_total_kobo=:total,file_sha256=:hash,imported_by=:actor,imported_at=now() WHERE cooperative_id=:cid AND id=:id"
            ),
            {
                "cid": cid,
                "id": run_id,
                "expected": expected_total_kobo,
                "total": report.file_total_kobo,
                "hash": digest,
                "actor": current.user_id,
            },
        )
        await append_audit(
            session,
            cooperative_id=cid,
            actor_user_id=current.user_id,
            role=role,
            action="payroll_imported",
            entity_type="payroll_run",
            entity_id=str(run_id),
            summary={"file_total_kobo": report.file_total_kobo},
            channel="api",
            ip=None,
        )
    return report


async def post_run(current: CurrentSession, run_id: uuid.UUID) -> int:
    _role(current, {"treasurer"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run = (
            (
                await session.execute(
                    text(
                        "SELECT * FROM payroll_runs WHERE cooperative_id=:cid AND id=:id FOR UPDATE"
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        if run is None or run["status"] not in {"imported", "posted"}:
            raise Conflict("Payroll run cannot be posted")
        totals = (
            (
                await session.execute(
                    text(
                        """SELECT coalesce(sum(expected_loan_kobo+expected_savings_kobo+
                          expected_dues_kobo) FILTER (WHERE classification NOT IN
                          ('duplicate','unknown_member')),0)::bigint AS components_kobo,
                          coalesce(sum(actual_kobo),0)::bigint AS actual_kobo
                          FROM payroll_items WHERE cooperative_id=:cid AND run_id=:run"""
                    ),
                    {"cid": cid, "run": run_id},
                )
            )
            .mappings()
            .one()
        )
        if int(totals["components_kobo"]) != int(run["expected_total_kobo"]):
            raise Conflict("Payroll components do not reconcile to the expected run total")
        if int(totals["actual_kobo"]) != int(run["file_total_kobo"]):
            raise Conflict("Payroll item amounts do not reconcile to the imported run total")
        items = (
            (
                await session.execute(
                    text(
                        "SELECT * FROM payroll_items WHERE cooperative_id=:cid AND run_id=:run AND classification IN ('matched','short','over') AND posted_repayment_id IS NULL ORDER BY id FOR UPDATE"
                    ),
                    {"cid": cid, "run": run_id},
                )
            )
            .mappings()
            .all()
        )
        posted = 0
        for item in items:
            loan = await session.scalar(
                text(
                    "SELECT id FROM loans WHERE cooperative_id=:cid AND member_id=:member AND status='repaying' AND NOT deductions_paused ORDER BY created_at LIMIT 1 FOR UPDATE"
                ),
                {"cid": cid, "member": item["member_id"]},
            )
            if loan is None:
                continue
            repayment_id = await post_repayment(
                current,
                loan_id=loan,
                member_id=item["member_id"],
                amount_kobo=int(item["actual_kobo"]),
                source="payroll",
                received_on=date(
                    int(run["period_year"]),
                    int(run["period_month"]),
                    monthrange(int(run["period_year"]), int(run["period_month"]))[1],
                ),
                reference=f"payroll:{run_id}:{item['id']}",
                idempotency_key=f"payroll:{run_id}:{item['member_id']}",
                ip=None,
                db_session=session,
            )
            await session.execute(
                text(
                    "UPDATE payroll_items SET posted_repayment_id=:repayment WHERE cooperative_id=:cid AND id=:id"
                ),
                {"cid": cid, "id": item["id"], "repayment": repayment_id},
            )
            posted += 1
        await session.execute(
            text(
                "UPDATE payroll_runs SET status='posted',posted_by=CASE WHEN status='imported' THEN :actor ELSE posted_by END,posted_at=now() WHERE cooperative_id=:cid AND id=:id AND status<>'signed_off'"
            ),
            {"cid": cid, "id": run_id, "actor": current.user_id},
        )
    return posted


async def sign_off(
    current: CurrentSession, run_id: uuid.UUID, bank_total_kobo: int
) -> tuple[bytes, bytes]:
    role = _role(current, {"admin"})
    cid = current.cooperative_id
    async with transaction(cid) as session:
        run = (
            (
                await session.execute(
                    text(
                        "SELECT * FROM payroll_runs WHERE cooperative_id=:cid AND id=:id FOR UPDATE"
                    ),
                    {"cid": cid, "id": run_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        if run is None:
            raise NotFound("Payroll run not found")
        poster_id = run["posted_by"] if run["posted_by"] is not None else run["imported_by"]
        if run["status"] == "signed_off" or poster_id == current.user_id:
            raise Conflict("A different administrator must sign off the payroll run")
        if run["status"] != "posted":
            raise Conflict("Payroll run must be posted before sign-off")
        if run["file_total_kobo"] != bank_total_kobo:
            raise Conflict("File total does not match bank receipt total")
        unresolved = await session.scalar(
            text(
                "SELECT EXISTS(SELECT 1 FROM payroll_items WHERE cooperative_id=:cid AND run_id=:run AND classification<>'matched' AND resolution_note IS NULL)"
            ),
            {"cid": cid, "run": run_id},
        )
        if unresolved:
            raise Conflict("Payroll exceptions require resolution notes before sign-off")
        items = (
            (
                await session.execute(
                    text(
                        "SELECT staff_no_raw,classification,actual_kobo,resolution_note FROM payroll_items WHERE cooperative_id=:cid AND run_id=:run ORDER BY staff_no_raw"
                    ),
                    {"cid": cid, "run": run_id},
                )
            )
            .mappings()
            .all()
        )
        counts = {
            key: sum(row["classification"] == key for row in items)
            for key in ("matched", "short", "over", "missing", "unknown_member", "duplicate")
        }
        export_items = []
        for row in items:
            item = dict(row)
            if item["resolution_note"] is not None:
                item["resolution_note"] = str(item["resolution_note"])[:500]
            export_items.append(item)
        report = {
            "run_id": str(run_id),
            "expected_total_kobo": run["expected_total_kobo"],
            "file_total_kobo": run["file_total_kobo"],
            "bank_receipt_total_kobo": bank_total_kobo,
            "variance_kobo": int(run["file_total_kobo"]) - bank_total_kobo,
            "counts": counts,
            "exceptions": [row for row in export_items if row["classification"] != "matched"],
        }
        await session.execute(
            text(
                "UPDATE payroll_runs SET status='signed_off',bank_receipt_total_kobo=:bank,signed_off_by=:actor,signed_off_at=now() WHERE cooperative_id=:cid AND id=:id"
            ),
            {"cid": cid, "id": run_id, "bank": bank_total_kobo, "actor": current.user_id},
        )
        await append_audit(
            session,
            cooperative_id=cid,
            actor_user_id=current.user_id,
            role=role,
            action="payroll_signed_off",
            entity_type="payroll_run",
            entity_id=str(run_id),
            summary={"bank_total_kobo": bank_total_kobo},
            channel="api",
            ip=None,
        )
    output = io.StringIO(newline="")
    writer = csv.DictWriter(
        output, fieldnames=["staff_no_raw", "classification", "actual_kobo", "resolution_note"]
    )
    writer.writeheader()
    writer.writerows(
        _sanitize_csv_row(
            row,
            {"staff_no_raw", "classification", "resolution_note"},
            {"resolution_note": 500},
        )
        for row in report["exceptions"]
    )
    return output.getvalue().encode(), json.dumps(report, sort_keys=True, default=str).encode()
